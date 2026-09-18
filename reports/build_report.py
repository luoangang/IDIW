from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports"
BUILD_DIR = OUT_DIR / "_build"
OUTPUT = OUT_DIR / "情报采集与持续监控系统阶段性设计汇报.docx"
BUILD_DIR.mkdir(parents=True, exist_ok=True)

FONT_PATH = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_BOLD_PATH = Path(r"C:\Windows\Fonts\msyhbd.ttc")

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "EAF2F8"
TEAL = "207567"
PALE_TEAL = "E7F2EF"
GOLD = "9C6500"
PALE_GOLD = "FFF2CC"
GRAY = "666666"
PALE_GRAY = "F2F2F2"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def font(size: int, bold: bool = False):
    path = FONT_BOLD_PATH if bold and FONT_BOLD_PATH.exists() else FONT_PATH
    return ImageFont.truetype(str(path), size)


def draw_centered(draw, box, text, fnt, fill):
    x1, y1, x2, y2 = box
    bounds = draw.textbbox((0, 0), text, font=fnt)
    width = bounds[2] - bounds[0]
    height = bounds[3] - bounds[1]
    draw.text(((x1 + x2 - width) / 2, (y1 + y2 - height) / 2 - 2), text, font=fnt, fill=fill)


def draw_arrow(draw, start, end, color="#566573", width=5, dashed=False):
    x1, y1 = start
    x2, y2 = end
    if dashed:
        steps = 12
        for i in range(0, steps, 2):
            a = i / steps
            b = min((i + 1) / steps, 1)
            draw.line((x1 + (x2 - x1) * a, y1 + (y2 - y1) * a,
                       x1 + (x2 - x1) * b, y1 + (y2 - y1) * b), fill=color, width=width)
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=width)
    import math
    angle = math.atan2(y2 - y1, x2 - x1)
    length = 16
    left = (x2 - length * math.cos(angle - 0.55), y2 - length * math.sin(angle - 0.55))
    right = (x2 - length * math.cos(angle + 0.55), y2 - length * math.sin(angle + 0.55))
    draw.polygon([(x2, y2), left, right], fill=color)


def create_flow_diagram(path: Path):
    img = Image.new("RGB", (2200, 720), "white")
    draw = ImageDraw.Draw(img)
    title_font = font(42, True)
    node_font = font(30, True)
    sub_font = font(22)
    draw.text((70, 42), "系统采集与持续监控流程", font=title_font, fill="#17365D")

    nodes = [
        (90, 200, 430, 480, "监控任务", "关键词  别名  周期\n来源选择  采集上限", "#EAF2F8", "#2F75B5"),
        (520, 200, 860, 480, "调度与执行", "手动或定时触发\n来源隔离  失败降级", "#F2F2F2", "#666666"),
        (950, 200, 1290, 480, "来源插件", "百科  RSS  搜索\n论文与专利适配器", "#FFF2CC", "#9C6500"),
        (1380, 200, 1720, 480, "标准化与去重", "统一字段协议\n身份哈希  内容哈希", "#E7F2EF", "#207567"),
        (1810, 200, 2150, 480, "模块化存储", "新闻  论文  专利\n运行记录  命中关系", "#EAF2F8", "#17365D"),
    ]
    for x1, y1, x2, y2, heading, body, bg, edge in nodes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=14, fill=bg, outline=edge, width=4)
        draw_centered(draw, (x1 + 18, y1 + 35, x2 - 18, y1 + 115), heading, node_font, edge)
        lines = body.split("\n")
        for idx, line in enumerate(lines):
            draw_centered(draw, (x1 + 15, y1 + 125 + idx * 58, x2 - 15, y1 + 180 + idx * 58), line, sub_font, "#333333")
    for i in range(len(nodes) - 1):
        draw_arrow(draw, (nodes[i][2] + 12, 340), (nodes[i + 1][0] - 12, 340), color="#7F8C8D", width=5)
    draw.text((92, 565), "当前已运行：news 模块与 4 个来源", font=font(24, True), fill="#17365D")
    draw.text((750, 565), "已预留：paper 与 patent 模块", font=font(24, True), fill="#207567")
    draw.text((1420, 565), "目标：每日或每周持续更新", font=font(24, True), fill="#9C6500")
    img.save(path, quality=95)


def create_er_diagram(path: Path):
    img = Image.new("RGB", (2600, 1500), "white")
    draw = ImageDraw.Draw(img)
    title_font = font(44, True)
    box_title_font = font(26, True)
    field_font = font(19)
    label_font = font(18, True)
    draw.text((70, 38), "数据库 ER 图  Schema v4", font=title_font, fill="#17365D")
    draw.text((70, 100), "实线表示数据库外键关系  虚线表示通过 module_name 和 resource_uid 维护的逻辑关系", font=font(21), fill="#555555")

    boxes = {
        "sources": (80, 220, 550, 555, "sources", ["PK  source_key", "module_name", "label", "adapter_name", "enabled  item_limit", "config_json  status"], "#EAF2F8", "#2F75B5"),
        "monitors": (760, 220, 1230, 585, "monitors", ["PK  id", "module_name", "query_text", "aliases_json", "schedule_type  time", "last_run_at  next_run_at"], "#EAF2F8", "#2F75B5"),
        "monitor_sources": (380, 690, 950, 1050, "monitor_sources", ["PK FK  monitor_id", "PK FK  source_key", "enabled  item_limit", "config_json", "cursor_json", "last_success_at  last_error"], "#FFF2CC", "#9C6500"),
        "monitor_runs": (1080, 690, 1580, 1050, "monitor_runs", ["PK  id", "FK  monitor_id", "module_name  query_text", "trigger_type  status", "source_results_json", "fetched  inserted  updated"], "#F2F2F2", "#666666"),
        "matches": (1730, 650, 2350, 1060, "monitor_resource_matches", ["PK FK  monitor_id", "PK  module_name", "PK  resource_uid", "first_run_id  last_run_id", "first_matched_at  last_matched_at", "match_count  relevance_score"], "#FFF2CC", "#9C6500"),
        "news": (80, 1160, 700, 1425, "news_resources", ["PK  id   UK  resource_uid", "source_key  resource_type", "title  url  summary  content", "identity_hash  canonical_hash", "first_run_id  last_run_id"], "#E7F2EF", "#207567"),
        "paper": (835, 1160, 1455, 1425, "paper_resources", ["PK  id   UK  resource_uid", "source_key  DOI  journal", "authors  keywords  affiliations", "identity_hash  canonical_hash", "first_run_id  last_run_id"], "#E7F2EF", "#207567"),
        "patent": (1590, 1160, 2210, 1425, "patent_resources", ["PK  id   UK  resource_uid", "source_key  patent_number", "applicants  inventors  class", "identity_hash  canonical_hash", "first_run_id  last_run_id"], "#E7F2EF", "#207567"),
        "schema": (2050, 220, 2510, 480, "schema_migrations", ["PK  version", "name", "applied_at"], "#F2F2F2", "#666666"),
    }

    for _, (x1, y1, x2, y2, heading, fields, bg, edge) in boxes.items():
        draw.rounded_rectangle((x1, y1, x2, y2), radius=12, fill=bg, outline=edge, width=4)
        draw.rectangle((x1, y1, x2, y1 + 60), fill=edge)
        draw.text((x1 + 20, y1 + 13), heading, font=box_title_font, fill="white")
        for i, field in enumerate(fields):
            draw.text((x1 + 22, y1 + 82 + i * 39), field, font=field_font, fill="#222222")

    # Core foreign keys.
    draw_arrow(draw, (550, 435), (380, 790), color="#2F75B5", width=5)
    draw.text((430, 610), "1 对多", font=label_font, fill="#2F75B5")
    draw_arrow(draw, (910, 585), (760, 690), color="#2F75B5", width=5)
    draw.text((810, 625), "1 对多", font=label_font, fill="#2F75B5")
    draw_arrow(draw, (1100, 585), (1240, 690), color="#2F75B5", width=5)
    draw.text((1130, 625), "1 对多", font=label_font, fill="#2F75B5")
    draw_arrow(draw, (1230, 450), (1880, 650), color="#9C6500", width=5)
    draw.text((1510, 490), "1 对多", font=label_font, fill="#9C6500")

    # Sources to module resource tables are source_key references at application level.
    draw_arrow(draw, (360, 555), (300, 1160), color="#207567", width=4, dashed=True)
    draw.arrow = None
    draw_arrow(draw, (450, 555), (1080, 1160), color="#207567", width=4, dashed=True)
    draw_arrow(draw, (520, 555), (1860, 1160), color="#207567", width=4, dashed=True)

    # Match table polymorphic links to three module tables.
    draw_arrow(draw, (1900, 1060), (540, 1160), color="#9C6500", width=4, dashed=True)
    draw_arrow(draw, (2020, 1060), (1210, 1160), color="#9C6500", width=4, dashed=True)
    draw_arrow(draw, (2140, 1060), (1900, 1160), color="#9C6500", width=4, dashed=True)
    draw.text((2240, 1250), "三选一逻辑关联", font=label_font, fill="#9C6500")
    img.save(path, quality=95)


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), LIGHT_BORDER)
        borders.append(node)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, heading in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(heading))
        set_run_font(r, size=9.5, bold=True, color=WHITE)
        if widths:
            cell.width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_shading(cell, "FFFFFF" if ridx % 2 == 0 else PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=9.2)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10.5, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = False
    set_run_font(p.add_run(f"•  {text}"), size=10.2)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run(text), size=9, color=GRAY)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    end = paragraph.add_run(" 页")
    set_run_font(end, size=9, color=GRAY)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    for name, size, before, after in (
        ("Title", 26, 0, 12),
        ("Heading 1", 17, 14, 8),
        ("Heading 2", 13, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("情报采集与持续监控系统  阶段性设计汇报"), size=8.5, color=GRAY)
    add_page_number(section.footer.paragraphs[0])


def build_document():
    flow_path = BUILD_DIR / "system_flow.png"
    er_path = BUILD_DIR / "database_er.png"
    create_flow_diagram(flow_path)
    create_er_diagram(er_path)

    doc = Document()
    configure_document(doc)

    # Cover.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("情报采集与持续监控系统")
    set_run_font(r, size=28, bold=True, color=BLACK)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_after = Pt(26)
    r = p2.add_run("阶段性设计汇报")
    set_run_font(r, size=22, bold=True, color=NAVY)

    add_body(doc, "本报告说明当前系统已经形成的模块化采集架构、精简后的数据库设计、已完成能力和下一阶段实施计划。当前结论是：新闻模块已经具备可运行的采集与统一存储基础，论文和专利模块已完成架构预留，数据库已收敛为可管理的 9 表结构。")
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    add_table(doc, ["当前版本", "已接入来源", "数据库结构", "当前重点"], [["Schema v4", "4 个新闻来源", "9 张物理表", "持续监控与模块扩展"]], [3.1, 4.0, 3.7, 5.0])
    doc.add_paragraph().paragraph_format.space_after = Pt(40)
    p = doc.add_paragraph()
    set_run_font(p.add_run("汇报日期  2026 年 9 月 18 日"), size=11, color=GRAY)
    p = doc.add_paragraph()
    set_run_font(p.add_run("项目阶段  原型验证与架构定型"), size=11, color=GRAY)

    doc.add_page_break()
    doc.add_heading("一 项目目标与当前结论", level=1)
    add_body(doc, "系统面向长期情报监控场景：用户提供关键词或主题，系统按天或按周持续访问多个公开来源，完成采集、标准化、去重、存储和结果展示。后续在新闻基础上增加论文和专利模块，并在统一数据基础上开展实体提取、主题归并和趋势分析。")
    doc.add_heading("当前阶段结论", level=2)
    for text in (
        "新闻模块已经接入百度百科、Wikipedia、Google 新闻 RSS 和 DDGS。",
        "来源插件按业务模块物理隔离，新增来源不需要修改核心采集引擎。",
        "新闻、论文和专利分别使用独立资源表，避免不同领域字段相互污染。",
        "数据库从历史上的 30 张物理表精简为 9 张，运行历史和既有新闻数据得到保留。",
        "论文和专利目前属于架构预留，尚未接入正式生产来源。",
    ):
        add_bullet(doc, text)

    doc.add_heading("当前能力边界", level=2)
    add_table(doc, ["模块", "当前状态", "已经具备", "尚待完成"], [
        ["新闻", "可运行", "四个来源、统一采集协议、去重入库、前端展示", "定时监控管理、更多稳定来源、内容分析"],
        ["论文", "已预留", "模块注册、目录隔离、专业字段表", "中文论文适配器、增量同步、DOI 归一化"],
        ["专利", "已预留", "模块注册、目录隔离、专业字段表", "专利数据源、法律状态更新、分类检索"],
    ], [2.2, 2.4, 6.3, 6.0])

    doc.add_page_break()
    doc.add_heading("二 系统总体架构", level=1)
    add_body(doc, "系统采用模块、来源、执行、存储和控制五层结构。核心原则是让采集器只负责访问外部来源并返回标准数据，由统一引擎负责运行状态、失败隔离和持久化。")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(flow_path), width=Inches(6.45))
    add_caption(doc, "图 1  系统采集与持续监控流程")
    add_table(doc, ["层级", "主要职责", "当前实现"], [
        ["模块层", "定义新闻、论文、专利的能力边界和资源类型", "news  paper  patent"],
        ["来源层", "封装外部站点、RSS、搜索接口或浏览器自动化", "四个 news 插件"],
        ["执行层", "创建运行记录，依次执行来源，隔离单源失败", "统一 CollectionRequest"],
        ["存储层", "统一去重并写入对应模块资源表", "MySQL Schema v4"],
        ["控制层", "配置来源、提交任务、查看日志与结果", "Flask API 与管理面板"],
    ], [2.2, 8.0, 6.3])

    doc.add_heading("采集标准化协议", level=2)
    add_body(doc, "每个来源插件接收同一种采集请求，返回统一的 SourceItem。标题、链接、摘要、正文、作者、发布者和时间等公共字段直接入库；来源特有字段进入 metadata_json。采集源无法直接操作数据库，从而保持插件可替换、可测试。")

    doc.add_page_break()
    doc.add_heading("三 数据存储设计", level=1)
    add_body(doc, "数据库采用关系型设计，并允许少量 JSON 承载结构不固定的配置和元数据。固定、重要、需要查询或关联的数据使用普通字段；实体之间的关系使用关系表，不把 ID 列表隐藏在 JSON 中。")
    doc.add_heading("九张表的职责", level=2)
    add_table(doc, ["表", "类别", "职责"], [
        ["sources", "来源配置", "注册所有采集源及模块、适配器、启停、上限和配置"],
        ["monitors", "监控配置", "保存关键词、别名、每日或每周周期及下次执行时间"],
        ["monitor_sources", "关系与状态", "指定任务使用哪些来源，并保存任务级限制和增量游标"],
        ["monitor_runs", "执行记录", "保存每次手动或定时执行的状态、数量、错误和来源摘要"],
        ["news_resources", "业务数据", "保存新闻、百科和搜索结果"],
        ["paper_resources", "业务数据", "保存论文及 DOI、作者、期刊、卷期、基金等字段"],
        ["patent_resources", "业务数据", "保存专利号、申请人、发明人、分类和法律状态"],
        ["monitor_resource_matches", "业务关系", "保存监控任务与资源之间的多对多命中关系"],
        ["schema_migrations", "系统管理", "记录数据库结构版本，不参与日常业务"],
    ], [4.4, 2.8, 9.4])

    doc.add_heading("为什么新闻只使用一张资源表", level=2)
    add_body(doc, "百度百科、Wikipedia、Google 新闻 RSS 和 DDGS 虽然返回结构不同，但都可以归一为标题、链接、摘要、正文、来源、发布时间等公共字段。通过 source_key 区分来源，通过 resource_type 区分百科、新闻文章和搜索结果，少量来源专属字段放入 metadata_json。新增新闻来源时不需要新增数据表。")

    # Landscape ER page.
    section = doc.add_section(start_type=2)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    doc.add_heading("四 数据库 ER 图", level=1)
    add_body(doc, "核心外键关系集中在监控配置和运行记录中。三张模块资源表共用 resource_uid 规则；monitor_resource_matches 通过 module_name 与 resource_uid 指向其中一张资源表，这是为了保持三个模块物理隔离。")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(er_path), width=Inches(8.7))
    add_caption(doc, "图 2  Schema v4 数据库实体关系图")

    # Return to portrait.
    section = doc.add_section(start_type=2)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    doc.add_heading("五 数据关系与去重策略", level=1)
    doc.add_heading("核心关系", level=2)
    add_table(doc, ["关系", "含义", "实现方式"], [
        ["monitors 到 sources", "一个任务选择多个来源，一个来源服务多个任务", "monitor_sources 多对多关系表"],
        ["monitors 到 runs", "一个任务产生多次执行记录", "monitor_runs.monitor_id 外键"],
        ["sources 到 resources", "一个来源产生多条模块资源", "各资源表 source_key"],
        ["monitors 到 resources", "一个资源可以被多个监控主题命中", "monitor_resource_matches"],
    ], [4.1, 6.8, 5.8])

    doc.add_heading("资源身份", level=2)
    add_body(doc, "资源身份不包含用户输入的搜索词。这样 F-22、F22 或猛禽战斗机等不同关键词再次发现同一来源记录时，会更新已有资源而不是重复插入。")
    for text in (
        "identity_hash：来源内部稳定身份，优先使用来源 ID、DOI、专利号，否则使用 URL。",
        "canonical_hash：识别不同来源可能指向的同一内容，但当前不强制合并。",
        "content_hash：判断标题、摘要、正文或元数据是否发生变化。",
        "resource_uid：跨运行稳定的资源编号，用于监控命中关系。",
    ):
        p = add_bullet(doc, text)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15

    doc.add_heading("JSON 的使用边界", level=2)
    add_table(doc, ["可以使用 JSON", "必须使用关系字段或关系表"], [
        ["来源特有配置和支持模式", "来源、任务和资源的主键"],
        ["分页游标和同步状态", "任务选择来源的多对多关系"],
        ["运行中各来源的结果摘要", "任务与资源的多对多命中关系"],
        ["来源特有元数据和列表字段", "需要频繁筛选、排序和统计的核心字段"],
    ], [8.3, 8.3])

    doc.add_heading("六 当前实现与验证结果", level=1)
    add_table(doc, ["项目", "结果", "说明"], [
        ["数据库迁移", "通过", "从 30 张物理表收敛到 9 张，迁移前已生成完整快照"],
        ["历史数据", "保留", "迁移验证时保留 108 条新闻资源和 5 次历史运行"],
        ["运行摘要", "保留", "20 条来源运行明细已合并进对应 monitor_runs"],
        ["来源插件", "正常", "四个 news 插件完成注册并由统一引擎调用"],
        ["数据库链路", "通过", "游标、入库、去重、运行摘要和监控命中关系均完成验证"],
        ["自动化测试", "通过", "当前 9 项测试全部通过"],
        ["管理面板", "可访问", "本地地址为 http://127.0.0.1:8080"],
    ], [3.4, 2.5, 10.8])

    doc.add_heading("现阶段限制", level=2)
    for text in (
        "持续监控的数据结构已经建立，但每日和每周调度管理界面尚未完整实现。",
        "论文和专利模块目前没有正式采集源，资源表仍为空。",
        "免费公开来源没有稳定性承诺，可能受到限流、反爬、页面结构变化或网络环境影响。",
        "跨来源的同一事件目前只生成 canonical_hash，尚未进行自动聚类和主题合并。",
        "AI 实体提取、相关性分析和自动报告属于下一阶段能力。",
    ):
        add_bullet(doc, text)

    doc.add_page_break()
    doc.add_heading("七 下一阶段计划", level=1)
    add_table(doc, ["阶段", "主要任务", "交付结果"], [
        ["第一阶段", "完成 monitors 和 monitor_sources 管理界面，接入每日和每周调度", "用户可创建、暂停和立即执行持续监控任务"],
        ["第二阶段", "接入中文论文 Agent 和确定性论文适配器", "paper 模块形成可用的检索与增量同步链路"],
        ["第三阶段", "选择专利公开数据源并实现专业字段映射", "patent 模块具备申请号、分类和法律状态采集"],
        ["第四阶段", "增加实体提取、别名归并、事件聚类和趋势分析", "由资料采集升级为情报分析与报告生成"],
    ], [2.6, 8.1, 6.0])

    doc.add_heading("八 汇报结论", level=1)
    add_body(doc, "当前系统已经完成从一次性爬虫 Demo 到模块化采集平台基础架构的转变。新闻模块验证了来源插件、统一请求、失败隔离、去重入库和前端展示的完整链路；精简后的 9 表数据库能够支持持续监控，同时保持结构清晰。下一步应优先完成调度与监控管理，再逐步接入论文和专利来源，避免在业务尚未验证前继续扩张数据库复杂度。")

    doc.add_heading("建议汇报顺序", level=2)
    for text in (
        "先说明业务目标：围绕关键词持续收集新闻、论文和专利。",
        "再展示当前成果：新闻模块可运行，四个来源已接入。",
        "用系统流程图解释插件化和模块隔离。",
        "用 ER 图解释为什么最终保留 9 张表，以及 JSON 只承担扩展信息。",
        "最后明确边界和下一阶段计划，避免把预留模块表述为已完成能力。",
    ):
        add_bullet(doc, text)

    doc.core_properties.title = "情报采集与持续监控系统阶段性设计汇报"
    doc.core_properties.subject = "系统架构 数据库设计 ER图 阶段性成果"
    doc.core_properties.author = "项目组"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
